"""
SlydAI Bot - PPTX Slide Deck Generator
Generates native PowerPoint (.pptx) files styled dynamically according to the selected theme palette, fonts, and balanced 2-column premium layouts.
"""

import asyncio
import logging
import re
from pathlib import Path
from typing import Optional, Any

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR, MSO_AUTO_SIZE
from pptx.enum.shapes import MSO_SHAPE

from config import PROMPTS_DIR

logger = logging.getLogger(__name__)

TEMP_PPTX_DIR = Path(__file__).parent / "temp_pptx"
TEMP_PPTX_DIR.mkdir(parents=True, exist_ok=True)


def _clean_text(text: str) -> str:
    """Strip HTML tags (e.g. <span>, <div>, CSS styles) and unescape HTML entities."""
    if not text:
        return ""
    # Strip HTML tags
    cleaned = re.sub(r"<[^>]+>", "", str(text))
    # Replace common HTML entities
    cleaned = (
        cleaned.replace("&amp;", "&")
        .replace("&quot;", '"')
        .replace("&lt;", "<")
        .replace("&gt;", ">")
        .replace("&#39;", "'")
        .replace("&nbsp;", " ")
    )
    return cleaned.strip()


def _hex_to_rgb(hex_str: str, default: RGBColor = RGBColor(30, 41, 59)) -> RGBColor:
    """Convert hex string (e.g. '#0B0B16' or '0B0B16') to RGBColor."""
    if not hex_str:
        return default
    hex_clean = hex_str.strip().lstrip("#")
    if len(hex_clean) == 3:
        hex_clean = "".join([c * 2 for c in hex_clean])
    if len(hex_clean) < 6:
        return default
    try:
        r = int(hex_clean[0:2], 16)
        g = int(hex_clean[2:4], 16)
        b = int(hex_clean[4:6], 16)
        return RGBColor(r, g, b)
    except Exception:
        return default


def _is_color_dark(hex_str: str) -> bool:
    """Determine if a hex color is dark or light."""
    hex_clean = hex_str.strip().lstrip("#")
    if len(hex_clean) == 3:
        hex_clean = "".join([c * 2 for c in hex_clean])
    if len(hex_clean) < 6:
        return True
    try:
        r = int(hex_clean[0:2], 16)
        g = int(hex_clean[2:4], 16)
        b = int(hex_clean[4:6], 16)
        # Perceived brightness formula
        brightness = (r * 299 + g * 587 + b * 114) / 1000
        return brightness < 128
    except Exception:
        return True


def parse_theme_specs(theme_id: str) -> dict:
    """Extract palette, fonts, and style details from theme README.md or PROMPT.md."""
    readme_path = PROMPTS_DIR / theme_id / "README.md"
    prompt_path = PROMPTS_DIR / theme_id / "PROMPT.md"
    
    content = ""
    if readme_path.exists():
        content += readme_path.read_text(encoding="utf-8")
    if prompt_path.exists():
        content += "\n" + prompt_path.read_text(encoding="utf-8")

    # Defaults
    specs = {
        "theme_id": theme_id,
        "bg": "#0B0B16",
        "surface": "#15152A",
        "border": "#2A2A45",
        "primary": "#7C5CFF",
        "secondary": "#36E0D0",
        "heading_color": "#F4F3FF",
        "body_color": "#C3C0DE",
        "muted_color": "#807CA6",
        "heading_font": "Sora",
        "body_font": "Inter",
        "is_dark": True,
    }

    if not content:
        return specs

    # 1. Extract Palette table if present in README
    for role, key in [
        ("Background", "bg"),
        ("Surface / panel", "surface"),
        ("Border", "border"),
        ("Primary accent", "primary"),
        ("Heading text", "heading_color"),
        ("Body text", "body_color"),
        ("Muted text", "muted_color"),
    ]:
        m = re.search(rf"\|\s*{re.escape(role)}\s*\|\s*`?(#[0-9A-Fa-f]{{3,8}})`?\s*\|", content, re.IGNORECASE)
        if m and m.group(1):
            specs[key] = m.group(1).strip()

    # 2. Extract colors from text descriptions if not found in table
    if specs["bg"] == "#0B0B16":
        bg_match = re.search(r"Background:\s*(?:[^#\n\r]*?)`?(#[0-9A-Fa-f]{3,8})`?", content, re.IGNORECASE)
        if bg_match:
            specs["bg"] = bg_match.group(1).strip()

    if specs["primary"] == "#7C5CFF":
        pri_match = re.search(r"(?:Primary|accent|accent is)\s*(?:[^#\n\r]*?)`?(#[0-9A-Fa-f]{3,8})`?", content, re.IGNORECASE)
        if pri_match:
            specs["primary"] = pri_match.group(1).strip()

    # 3. Detect Mode (Dark vs Light)
    if "Mode: Light" in content or "Mode:** Light" in content:
        specs["is_dark"] = False
    elif "Mode: Dark" in content or "Mode:** Dark" in content:
        specs["is_dark"] = True
    else:
        specs["is_dark"] = _is_color_dark(specs["bg"])

    # If it's a light theme, make sure text colors are appropriately contrasting
    if not specs["is_dark"]:
        if specs["bg"] == "#0B0B16":
            specs["bg"] = "#FFFFFF"
        if specs["surface"] == "#15152A":
            specs["surface"] = "#F8FAFC"
        if specs["border"] == "#2A2A45":
            specs["border"] = "#E2E8F0"
        if specs["heading_color"] == "#F4F3FF":
            specs["heading_color"] = "#0F172A"
        if specs["body_color"] == "#C3C0DE":
            specs["body_color"] = "#334155"
        if specs["muted_color"] == "#807CA6":
            specs["muted_color"] = "#64748B"
        if specs["primary"] == "#7C5CFF":
            specs["primary"] = "#2563EB"

    # 4. Extract Fonts
    hf_match = re.search(r"-\s*\*\*([^*\n\r]+)\*\*\s*\(heading", content, re.IGNORECASE)
    if hf_match and hf_match.group(1):
        specs["heading_font"] = hf_match.group(1).strip()
    else:
        hf2 = re.search(r"headlines in ['\"]?([A-Za-z\s]+)['\"]?", content, re.IGNORECASE)
        if hf2 and hf2.group(1):
            specs["heading_font"] = hf2.group(1).strip()

    bf_match = re.search(r"-\s*\*\*([^*\n\r]+)\*\*\s*\((?:supporting|body)", content, re.IGNORECASE)
    if bf_match and bf_match.group(1):
        specs["body_font"] = bf_match.group(1).strip()
    else:
        bf2 = re.search(r"body (?:and labels )?in ['\"]?([A-Za-z\s]+)['\"]?", content, re.IGNORECASE)
        if bf2 and bf2.group(1):
            specs["body_font"] = bf2.group(1).strip()

    return specs


def _set_box_margins(tf, top_pt=8, bottom_pt=8, left_pt=12, right_pt=12):
    """Set inner margins for a shape text frame."""
    tf.margin_top = Pt(top_pt)
    tf.margin_bottom = Pt(bottom_pt)
    tf.margin_left = Pt(left_pt)
    tf.margin_right = Pt(right_pt)


def _presentation_items(items: list[dict], max_body_words: int = 32) -> list[dict]:
    """Remove repeated points and keep each card readable at presentation size."""
    result = []
    seen = set()

    for item in items:
        lead = _clean_text(item.get("lead", ""))
        body_words = _clean_text(item.get("body", "")).split()
        body = " ".join(body_words[:max_body_words])
        if len(body_words) > max_body_words:
            body += "…"

        fingerprint = re.sub(r"\W+", "", f"{lead} {body}".lower())
        if not fingerprint or fingerprint in seen:
            continue
        seen.add(fingerprint)
        result.append({"lead": lead, "body": body})

    return result


def _prepare_card_text_frame(tf, top_pt=12, bottom_pt=10, left_pt=20, right_pt=18):
    """Apply safe text settings so card content stays within its bounds."""
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
    tf.vertical_anchor = MSO_ANCHOR.TOP
    _set_box_margins(tf, top_pt, bottom_pt, left_pt, right_pt)


def _add_styled_bullet_paragraph(
    tf,
    lead: str,
    body: str,
    font_name: str,
    lead_color: RGBColor,
    body_color: RGBColor,
    font_size: Pt = Pt(14),
    space_after: Pt = Pt(10),
    bullet_symbol: str = "• ",
):
    """Add a beautifully formatted paragraph with bold lead and readable body."""
    p = tf.add_paragraph()
    p.space_after = space_after
    
    clean_lead = _clean_text(lead)
    clean_body = _clean_text(body)

    if clean_lead:
        r_bullet = p.add_run()
        r_bullet.text = bullet_symbol
        r_bullet.font.name = font_name
        r_bullet.font.size = font_size
        r_bullet.font.bold = True
        r_bullet.font.color.rgb = lead_color

        r_lead = p.add_run()
        r_lead.text = f"{clean_lead}: "
        r_lead.font.name = font_name
        r_lead.font.size = Pt(font_size.pt + 0.5)
        r_lead.font.bold = True
        r_lead.font.color.rgb = lead_color

    r_body = p.add_run()
    if not clean_lead:
        r_body.text = f"{bullet_symbol}{clean_body}"
    else:
        r_body.text = clean_body

    r_body.font.name = font_name
    r_body.font.size = font_size
    r_body.font.color.rgb = body_color


def create_presentation_file(slides: list[dict], theme_id: str, topic: str = "") -> Path:
    """Create a fully styled .pptx presentation file based on theme specs with 2-column layouts and zero watermarks."""
    theme_specs = parse_theme_specs(theme_id)
    clean_topic = _clean_text(topic).strip()
    
    prs = Presentation()
    # 16:9 widescreen dimensions (13.333 x 7.5 inches)
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_slide_layout = prs.slide_layouts[6]

    bg_rgb = _hex_to_rgb(theme_specs.get("bg", "#0B0B16"))
    surface_rgb = _hex_to_rgb(theme_specs.get("surface", "#15152A"))
    border_rgb = _hex_to_rgb(theme_specs.get("border", "#2A2A45"))
    primary_rgb = _hex_to_rgb(theme_specs.get("primary", "#7C5CFF"))
    heading_rgb = _hex_to_rgb(theme_specs.get("heading_color", "#F4F3FF"))
    body_rgb = _hex_to_rgb(theme_specs.get("body_color", "#C3C0DE"))
    muted_rgb = _hex_to_rgb(theme_specs.get("muted_color", "#807CA6"))

    heading_font = theme_specs.get("heading_font", "Sora")
    body_font = theme_specs.get("body_font", "Inter")

    total_slides = len(slides)

    for slide_data in slides:
        idx = slide_data.get("index", 1)
        
        # Strictly ensure Slide 1 title is user topic
        if idx == 1:
            title = clean_topic if clean_topic else _clean_text(slide_data.get("title", f"Slide {idx}"))
        else:
            title = _clean_text(slide_data.get("title", f"Slide {idx}"))

        subtitle = _clean_text(slide_data.get("subtitle", ""))
        bullets = slide_data.get("bullets", [])
        structured_bullets = slide_data.get("structured_bullets", [])

        # Fallback for structured bullets if not present
        if not structured_bullets:
            for b in bullets:
                clean_b = _clean_text(b)
                m_lead = re.match(r"^\*\*(.*?)\*\*:?\s*(.*)$", clean_b)
                if m_lead:
                    structured_bullets.append({"lead": _clean_text(m_lead.group(1)), "body": _clean_text(m_lead.group(2))})
                else:
                    m_c = re.match(r"^([^:]{3,35}):\s*(.*)$", clean_b)
                    if m_c:
                        structured_bullets.append({"lead": _clean_text(m_c.group(1)), "body": _clean_text(m_c.group(2))})
                    else:
                        structured_bullets.append({"lead": "", "body": clean_b})

        # Cards use concise, unique points. This prevents repeated content in
        # two columns and avoids text running outside a card.
        structured_bullets = _presentation_items(structured_bullets)

        slide = prs.slides.add_slide(blank_slide_layout)

        # ── 1. Full Canvas Background ──
        bg_shape = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5)
        )
        bg_shape.fill.solid()
        bg_shape.fill.fore_color.rgb = bg_rgb
        bg_shape.line.fill.background()

        # ── 2. Top Unity Accent Bar ──
        top_bar = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(0.08)
        )
        top_bar.fill.solid()
        top_bar.fill.fore_color.rgb = primary_rgb
        top_bar.line.fill.background()

        # ── 3. Layout Router ──
        
        # ═══ COVER SLIDE (Slide 1) ═══
        if idx == 1:
            # Main Hero Card centered gracefully
            main_card = slide.shapes.add_shape(
                MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.0), Inches(1.3), Inches(11.333), Inches(4.85)
            )
            main_card.fill.solid()
            main_card.fill.fore_color.rgb = surface_rgb
            main_card.line.color.rgb = border_rgb
            main_card.line.width = Pt(1.5)

            # Left vertical accent line inside hero card
            side_line = slide.shapes.add_shape(
                MSO_SHAPE.RECTANGLE, Inches(1.0), Inches(1.3), Inches(0.14), Inches(4.85)
            )
            side_line.fill.solid()
            side_line.fill.fore_color.rgb = primary_rgb
            side_line.line.fill.background()

            tf = main_card.text_frame
            tf.word_wrap = True
            _set_box_margins(tf, top_pt=40, bottom_pt=30, left_pt=45, right_pt=45)

            # Main Title (Strictly user given topic)
            p_title = tf.paragraphs[0]
            p_title.text = title
            p_title.alignment = PP_ALIGN.LEFT
            p_title.font.name = heading_font
            p_title.font.size = Pt(44) if len(title) < 40 else (Pt(36) if len(title) < 75 else Pt(30))
            p_title.font.bold = True
            p_title.font.color.rgb = heading_rgb
            p_title.space_after = Pt(20)

            # Subtitle
            sub_text = subtitle if subtitle else "Professional tahliliy taqdimot"
            p_sub = tf.add_paragraph()
            p_sub.text = sub_text
            p_sub.alignment = PP_ALIGN.LEFT
            p_sub.font.name = body_font
            p_sub.font.size = Pt(20)
            p_sub.font.color.rgb = body_rgb
            p_sub.space_after = Pt(24)

            # Decorative Accent Bottom Divider
            p_div = tf.add_paragraph()
            p_div.text = f"━━━━━━━━━━━━━━━━━━━━━━━"
            p_div.font.name = body_font
            p_div.font.size = Pt(12)
            p_div.font.color.rgb = primary_rgb

        # ═══ AGENDA SLIDE (Slide 2) — 2-COLUMN GRID ═══
        elif idx == 2:
            # Tag Chip
            tag = slide.shapes.add_shape(
                MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(0.5), Inches(2.8), Inches(0.38)
            )
            tag.fill.solid()
            tag.fill.fore_color.rgb = surface_rgb
            tag.line.color.rgb = primary_rgb
            tag.line.width = Pt(1)
            tf_t = tag.text_frame
            _set_box_margins(tf_t, top_pt=3, bottom_pt=3, left_pt=8, right_pt=8)
            p_t = tf_t.paragraphs[0]
            p_t.text = "REJA VA MUNDARIJA"
            p_t.alignment = PP_ALIGN.CENTER
            p_t.font.name = body_font
            p_t.font.size = Pt(11)
            p_t.font.bold = True
            p_t.font.color.rgb = primary_rgb

            # Title
            tx_title = slide.shapes.add_textbox(Inches(0.8), Inches(0.92), Inches(11.733), Inches(0.95))
            tf = tx_title.text_frame
            tf.word_wrap = True
            _set_box_margins(tf, top_pt=0, bottom_pt=0, left_pt=0, right_pt=0)
            p_title = tf.paragraphs[0]
            p_title.text = title if title else "Taqdimotning Asosiy Yo'nalishlari"
            p_title.font.name = heading_font
            p_title.font.size = Pt(32)
            p_title.font.bold = True
            p_title.font.color.rgb = heading_rgb

            # 2-Column Grid for Agenda Items
            items = structured_bullets[:6] if structured_bullets else [
                {"lead": "Kirish va Kontekst", "body": "Mavzuning dolzarbligi va umumiy maqsadi"},
                {"lead": "Asosiy Tahlil", "body": "Tushunchalar, mexanizmlar va metodologiya"},
                {"lead": "Amaliy Tadbiq", "body": "Haqiqiy misollar, natijalar va strategiyalar"},
                {"lead": "Xulosa va Tavsiyalar", "body": "Yakuniy xulosalar va kelgusi qadamlar"},
            ]

            col_w = Inches(5.72)
            col_gap = Inches(0.293)
            
            # Divide into rows and 2 columns
            total_items = len(items)
            rows = 2 if total_items <= 4 else 3
            card_h = Inches(4.65 / rows - 0.15)
            
            for a_i, item in enumerate(items):
                col_idx = a_i % 2
                row_idx = a_i // 2
                
                left_pos = Inches(0.8) if col_idx == 0 else Inches(0.8 + 5.72 + 0.293)
                top_pos = Inches(2.05) + row_idx * (card_h + Inches(0.18))
                
                # Card Background
                card = slide.shapes.add_shape(
                    MSO_SHAPE.ROUNDED_RECTANGLE, left_pos, top_pos, col_w, card_h
                )
                card.fill.solid()
                card.fill.fore_color.rgb = surface_rgb
                card.line.color.rgb = border_rgb
                card.line.width = Pt(1.5)

                # Top Accent strip on card
                card_top = slide.shapes.add_shape(
                    MSO_SHAPE.RECTANGLE, left_pos, top_pos, col_w, Inches(0.06)
                )
                card_top.fill.solid()
                card_top.fill.fore_color.rgb = primary_rgb
                card_top.line.fill.background()

                # Number indicator pill
                num_pill = slide.shapes.add_shape(
                    MSO_SHAPE.ROUNDED_RECTANGLE, left_pos + Inches(0.18), top_pos + Inches(0.18), Inches(0.75), Inches(0.42)
                )
                num_pill.fill.solid()
                num_pill.fill.fore_color.rgb = bg_rgb
                num_pill.line.color.rgb = primary_rgb
                num_pill.line.width = Pt(1.5)
                num_tf = num_pill.text_frame
                _set_box_margins(num_tf, top_pt=2, bottom_pt=2, left_pt=2, right_pt=2)
                num_p = num_tf.paragraphs[0]
                num_p.text = f"0{a_i+1}"
                num_p.alignment = PP_ALIGN.CENTER
                num_p.font.name = body_font
                num_p.font.size = Pt(13.5)
                num_p.font.bold = True
                num_p.font.color.rgb = primary_rgb

                # Content Text Frame inside Card
                card_tf = card.text_frame
                _prepare_card_text_frame(card_tf, top_pt=14, bottom_pt=10, left_pt=78, right_pt=16)
                
                lead_t = _clean_text(item.get("lead", "").strip()) or f"Yo'nalish 0{a_i+1}"
                body_t = _clean_text(item.get("body", "").strip())
                
                p_lead = card_tf.paragraphs[0]
                p_lead.text = lead_t
                p_lead.font.name = heading_font
                p_lead.font.size = Pt(16)
                p_lead.font.bold = True
                p_lead.font.color.rgb = heading_rgb
                p_lead.space_after = Pt(4)

                if body_t:
                    p_body = card_tf.add_paragraph()
                    p_body.text = body_t
                    p_body.font.name = body_font
                    p_body.font.size = Pt(13)
                    p_body.font.color.rgb = body_rgb

        # ═══ SUMMARY / CONCLUSION SLIDE (Final Slide) — 2-COLUMN STRATEGIC CARDS ═══
        elif idx == total_slides:
            # Badge
            badge = slide.shapes.add_shape(
                MSO_SHAPE.ROUNDED_RECTANGLE, Inches(4.1), Inches(0.5), Inches(5.133), Inches(0.38)
            )
            badge.fill.solid()
            badge.fill.fore_color.rgb = surface_rgb
            badge.line.color.rgb = primary_rgb
            badge.line.width = Pt(1)
            tf_b = badge.text_frame
            _set_box_margins(tf_b, top_pt=3, bottom_pt=3, left_pt=8, right_pt=8)
            p_b = tf_b.paragraphs[0]
            p_b.text = "XULOSA VA STRATEGIK TAVSIYALAR"
            p_b.alignment = PP_ALIGN.CENTER
            p_b.font.name = body_font
            p_b.font.size = Pt(11)
            p_b.font.bold = True
            p_b.font.color.rgb = primary_rgb

            # Title
            tx_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.92), Inches(11.733), Inches(1.15))
            tf = tx_box.text_frame
            tf.word_wrap = True
            _set_box_margins(tf, top_pt=0, bottom_pt=0, left_pt=0, right_pt=0)
            p_title = tf.paragraphs[0]
            p_title.text = title if title else "Yakuniy Xulosa va Harakatlar Rejasi"
            p_title.alignment = PP_ALIGN.CENTER
            p_title.font.name = heading_font
            p_title.font.size = Pt(30)
            p_title.font.bold = True
            p_title.font.color.rgb = heading_rgb

            if subtitle:
                p_sub = tf.add_paragraph()
                p_sub.text = subtitle
                p_sub.alignment = PP_ALIGN.CENTER
                p_sub.font.name = body_font
                p_sub.font.size = Pt(15.5)
                p_sub.font.color.rgb = body_rgb
                p_sub.space_before = Pt(4)

            # 2 Large Column Cards for Strategy & Action Plan
            col_w = Inches(5.72)
            col_h = Inches(4.65)
            
            half = (len(structured_bullets) + 1) // 2
            left_items = structured_bullets[:half] if structured_bullets else [{"lead": "Asosiy Xulosalar", "body": "Mavzu bo'yicha shakllangan strategik xulosalar va tahliliy natijalar."}]
            right_items = structured_bullets[half:] if len(structured_bullets) > half else [{"lead": "Amaliy Tavsiyalar", "body": "Tizimni muvaffaqiyatli joriy etish va samaradorlikni oshirish bo'yicha muhim qadamlar."}]

            # Left Conclusion Card
            card_l = slide.shapes.add_shape(
                MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(2.05), col_w, col_h
            )
            card_l.fill.solid()
            card_l.fill.fore_color.rgb = surface_rgb
            card_l.line.color.rgb = border_rgb
            card_l.line.width = Pt(1.5)

            card_lt = slide.shapes.add_shape(
                MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(2.05), col_w, Inches(0.06)
            )
            card_lt.fill.solid()
            card_lt.fill.fore_color.rgb = primary_rgb
            card_lt.line.fill.background()

            tf_l = card_l.text_frame
            _prepare_card_text_frame(tf_l, top_pt=20, bottom_pt=14, left_pt=22, right_pt=22)

            p_lh = tf_l.paragraphs[0]
            p_lh.text = "01. Strategik Xulosa va Tahlil"
            p_lh.font.name = heading_font
            p_lh.font.size = Pt(17)
            p_lh.font.bold = True
            p_lh.font.color.rgb = primary_rgb
            p_lh.space_after = Pt(14)

            for item in left_items:
                _add_styled_bullet_paragraph(
                    tf_l,
                    item.get("lead", ""),
                    item.get("body", ""),
                    body_font,
                    heading_rgb,
                    body_rgb,
                    font_size=Pt(14),
                    space_after=Pt(12),
                )

            # Right Action Card
            card_r = slide.shapes.add_shape(
                MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.813), Inches(2.05), col_w, col_h
            )
            card_r.fill.solid()
            card_r.fill.fore_color.rgb = surface_rgb
            card_r.line.color.rgb = border_rgb
            card_r.line.width = Pt(1.5)

            card_rt = slide.shapes.add_shape(
                MSO_SHAPE.RECTANGLE, Inches(6.813), Inches(2.05), col_w, Inches(0.06)
            )
            card_rt.fill.solid()
            card_rt.fill.fore_color.rgb = primary_rgb
            card_rt.line.fill.background()

            tf_r = card_r.text_frame
            _prepare_card_text_frame(tf_r, top_pt=20, bottom_pt=14, left_pt=22, right_pt=22)

            p_rh = tf_r.paragraphs[0]
            p_rh.text = "02. Amaliy Harakatlar Rejasi"
            p_rh.font.name = heading_font
            p_rh.font.size = Pt(17)
            p_rh.font.bold = True
            p_rh.font.color.rgb = primary_rgb
            p_rh.space_after = Pt(14)

            for item in right_items:
                _add_styled_bullet_paragraph(
                    tf_r,
                    item.get("lead", ""),
                    item.get("body", ""),
                    body_font,
                    heading_rgb,
                    body_rgb,
                    font_size=Pt(14),
                    space_after=Pt(12),
                )

        # ═══ RICH CONTENT SLIDES (Slide 3 to Total - 1) — 2-COLUMN LAYOUTS ═══
        else:
            # Header Tag
            tag = slide.shapes.add_shape(
                MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(0.5), Inches(2.6), Inches(0.38)
            )
            tag.fill.solid()
            tag.fill.fore_color.rgb = surface_rgb
            tag.line.color.rgb = border_rgb
            tag.line.width = Pt(1)
            tf_t = tag.text_frame
            _set_box_margins(tf_t, top_pt=3, bottom_pt=3, left_pt=8, right_pt=8)
            p_t = tf_t.paragraphs[0]
            p_t.text = f"SLAYD {idx:02d} / {total_slides:02d}"
            p_t.alignment = PP_ALIGN.CENTER
            p_t.font.name = body_font
            p_t.font.size = Pt(11)
            p_t.font.bold = True
            p_t.font.color.rgb = primary_rgb

            # Title & Subtitle Header
            tx_title = slide.shapes.add_textbox(Inches(0.8), Inches(0.92), Inches(11.733), Inches(1.1))
            tf = tx_title.text_frame
            tf.word_wrap = True
            _set_box_margins(tf, top_pt=0, bottom_pt=0, left_pt=0, right_pt=0)
            p_title = tf.paragraphs[0]
            p_title.text = title
            p_title.font.name = heading_font
            p_title.font.size = Pt(28) if len(title) < 55 else Pt(24)
            p_title.font.bold = True
            p_title.font.color.rgb = heading_rgb

            if subtitle:
                p_sub = tf.add_paragraph()
                p_sub.text = subtitle
                p_sub.font.name = body_font
                p_sub.font.size = Pt(15.5)
                p_sub.font.color.rgb = body_rgb
                p_sub.space_before = Pt(3)

            # Balanced 2-Column Layout Types (Alternating for visual beauty)
            layout_type = idx % 2

            # ── Layout 1: 2-Column Balanced Dual Cards ──
            if layout_type == 0:
                col_w = Inches(5.72)
                col_h = Inches(4.65)
                
                half = (len(structured_bullets) + 1) // 2
                left_items = structured_bullets[:half]
                # Do not mirror left-side content when there are too few points.
                right_items = structured_bullets[half:]

                # Left Column Box
                card_l = slide.shapes.add_shape(
                    MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(2.05), col_w, col_h
                )
                card_l.fill.solid()
                card_l.fill.fore_color.rgb = surface_rgb
                card_l.line.color.rgb = border_rgb
                card_l.line.width = Pt(1.5)

                card_lt = slide.shapes.add_shape(
                    MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(2.05), col_w, Inches(0.06)
                )
                card_lt.fill.solid()
                card_lt.fill.fore_color.rgb = primary_rgb
                card_lt.line.fill.background()

                tf_l = card_l.text_frame
                _prepare_card_text_frame(tf_l, top_pt=18, bottom_pt=14, left_pt=20, right_pt=20)
                
                p_lh = tf_l.paragraphs[0]
                p_lh.text = "Asosiy Tushunchalar va Tahlil"
                p_lh.font.name = heading_font
                p_lh.font.size = Pt(16)
                p_lh.font.bold = True
                p_lh.font.color.rgb = primary_rgb
                p_lh.space_after = Pt(12)

                for item in left_items:
                    _add_styled_bullet_paragraph(
                        tf_l,
                        item.get("lead", ""),
                        item.get("body", ""),
                        body_font,
                        heading_rgb,
                        body_rgb,
                        font_size=Pt(14),
                        space_after=Pt(12),
                    )

                # Right Column Box
                card_r = slide.shapes.add_shape(
                    MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.813), Inches(2.05), col_w, col_h
                )
                card_r.fill.solid()
                card_r.fill.fore_color.rgb = surface_rgb
                card_r.line.color.rgb = border_rgb
                card_r.line.width = Pt(1.5)

                card_rt = slide.shapes.add_shape(
                    MSO_SHAPE.RECTANGLE, Inches(6.813), Inches(2.05), col_w, Inches(0.06)
                )
                card_rt.fill.solid()
                card_rt.fill.fore_color.rgb = primary_rgb
                card_rt.line.fill.background()

                tf_r = card_r.text_frame
                _prepare_card_text_frame(tf_r, top_pt=18, bottom_pt=14, left_pt=20, right_pt=20)
                
                p_rh = tf_r.paragraphs[0]
                p_rh.text = "Amaliy Tatbiq va Strategik Natijalar"
                p_rh.font.name = heading_font
                p_rh.font.size = Pt(16)
                p_rh.font.bold = True
                p_rh.font.color.rgb = primary_rgb
                p_rh.space_after = Pt(12)

                for item in right_items:
                    _add_styled_bullet_paragraph(
                        tf_r,
                        item.get("lead", ""),
                        item.get("body", ""),
                        body_font,
                        heading_rgb,
                        body_rgb,
                        font_size=Pt(14),
                        space_after=Pt(12),
                    )

            # ── Layout 2: 2x2 Grid (4 distinct Cards in 2 Columns x 2 Rows) ──
            else:
                items = structured_bullets[:4] if structured_bullets else [
                    {"lead": "1-Yo'nalish", "body": "Asosiy jarayon va tahlil"},
                    {"lead": "2-Yo'nalish", "body": "Amaliy mexanizmlar va metodlar"},
                    {"lead": "3-Yo'nalish", "body": "Monitoring va samaradorlik"},
                    {"lead": "4-Yo'nalish", "body": "Kutilayotgan natijalar va o'sish"},
                ]
                
                col_w = Inches(5.72)
                card_h = Inches(2.23)
                
                for b_i, item in enumerate(items[:4]):
                    col_idx = b_i % 2
                    row_idx = b_i // 2
                    
                    left_pos = Inches(0.8) if col_idx == 0 else Inches(6.813)
                    top_pos = Inches(2.05) + row_idx * (card_h + Inches(0.18))

                    # Card Box
                    card = slide.shapes.add_shape(
                        MSO_SHAPE.ROUNDED_RECTANGLE, left_pos, top_pos, col_w, card_h
                    )
                    card.fill.solid()
                    card.fill.fore_color.rgb = surface_rgb
                    card.line.color.rgb = border_rgb
                    card.line.width = Pt(1.5)

                    # Card Left Accent strip
                    side_accent = slide.shapes.add_shape(
                        MSO_SHAPE.RECTANGLE, left_pos, top_pos, Inches(0.08), card_h
                    )
                    side_accent.fill.solid()
                    side_accent.fill.fore_color.rgb = primary_rgb
                    side_accent.line.fill.background()

                    card_tf = card.text_frame
                    _prepare_card_text_frame(card_tf, top_pt=12, bottom_pt=10, left_pt=20, right_pt=18)

                    lead_t = _clean_text(item.get("lead", "").strip()) or f"Pillar 0{b_i+1}"
                    body_t = _clean_text(item.get("body", "").strip())

                    p_lead = card_tf.paragraphs[0]
                    p_lead.text = f"✦  {lead_t}"
                    p_lead.font.name = heading_font
                    p_lead.font.size = Pt(15.5)
                    p_lead.font.bold = True
                    p_lead.font.color.rgb = primary_rgb
                    p_lead.space_after = Pt(6)

                    if body_t:
                        p_body = card_tf.add_paragraph()
                        p_body.text = body_t
                        p_body.font.name = body_font
                        p_body.font.size = Pt(13.5)
                        p_body.font.color.rgb = body_rgb

        # ── 4. Slide Footer (Clean Page Indicator with Zero Watermarks) ──
        footer_tx = slide.shapes.add_textbox(Inches(0.8), Inches(6.92), Inches(11.733), Inches(0.35))
        f_tf = footer_tx.text_frame
        _set_box_margins(f_tf, top_pt=0, bottom_pt=0, left_pt=0, right_pt=0)
        f_p = f_tf.paragraphs[0]
        f_p.alignment = PP_ALIGN.RIGHT
        f_p.text = f"{idx} / {total_slides}"
        f_p.font.name = body_font
        f_p.font.size = Pt(11)
        f_p.font.color.rgb = muted_rgb

    # Save presentation file strictly using topic name
    safe_topic = re.sub(r"[^\w\s-]", "", clean_topic).strip()
    safe_topic = re.sub(r"[-\s]+", "_", safe_topic)[:60] or "Taqdimot"
    filename = f"{safe_topic}.pptx"
    out_path = TEMP_PPTX_DIR / filename
    prs.save(str(out_path))
    return out_path


async def generate_pptx_file(slides: list[dict], theme_id: str, topic: str = "") -> Path:
    """Async wrapper to build PPTX file in executor."""
    loop = asyncio.get_running_loop()
    return await loop.run_in_executor(None, create_presentation_file, slides, theme_id, topic)
